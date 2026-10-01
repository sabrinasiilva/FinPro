import sqlite3
from datetime import date
from io import BytesIO

import openpyxl

import app as finpro

HOJE = date.today().isoformat()


def entrar(client, usuario='ana', senha='senha123'):
    client.post('/register', data={'username': usuario, 'senha': senha})
    return client.post('/login', data={'username': usuario, 'senha': senha})


def consultar(sql, params=()):
    conn = sqlite3.connect(finpro.app.config['DATABASE'])
    linhas = conn.execute(sql, params).fetchall()
    conn.close()
    return linhas


def adicionar_investimento(client, ativo='ITSA4'):
    client.post('/investimentos', data={
        'tipo': 'Ações', 'ativo': ativo, 'quantidade': '10', 'valor_unitario': '9.50',
        'valor_atual': '10.00', 'data_compra': HOJE, 'descricao': '',
    })
    return consultar('SELECT id FROM investimentos WHERE ativo = ?', (ativo,))[0][0]


# ---------------------------
# Login e permissões
# ---------------------------
def test_login_com_senha_errada_nao_entra(client):
    entrar(client)
    client.get('/logout')
    resposta = client.post('/login', data={'username': 'ana', 'senha': 'errada'})
    assert resposta.status_code == 200
    assert 'Usuário ou senha incorretos' in resposta.get_data(as_text=True)


def test_paginas_exigem_login(client):
    for url in ['/dashboard', '/receitas', '/despesas', '/investimentos', '/gerar_relatorio']:
        resposta = client.get(url)
        assert resposta.status_code == 302, url
        assert '/login' in resposta.headers['Location'], url


def test_excluir_investimento_exige_login(app):
    ana = app.test_client()
    entrar(ana)
    investimento_id = adicionar_investimento(ana)

    visitante = app.test_client()
    resposta = visitante.post(f'/delete_investimento/{investimento_id}')

    assert resposta.status_code == 302
    assert consultar('SELECT COUNT(*) FROM investimentos')[0][0] == 1


def test_usuario_nao_apaga_investimento_de_outro(app):
    ana = app.test_client()
    entrar(ana, 'ana')
    investimento_id = adicionar_investimento(ana)

    bia = app.test_client()
    entrar(bia, 'bia')
    bia.post(f'/delete_investimento/{investimento_id}')

    assert consultar('SELECT COUNT(*) FROM investimentos')[0][0] == 1


def test_usuario_so_ve_os_proprios_lancamentos(app):
    ana = app.test_client()
    entrar(ana, 'ana')
    ana.post('/receitas', data={'valor': '100', 'categoria': 'Salário', 'descricao': 'Salário da Ana', 'data': HOJE})

    bia = app.test_client()
    entrar(bia, 'bia')
    assert 'Salário da Ana' not in bia.get('/receitas').get_data(as_text=True)


# ---------------------------
# Receitas e despesas
# ---------------------------
def test_editar_receita(client):
    entrar(client)
    client.post('/receitas', data={'valor': '100', 'categoria': 'Salário', 'descricao': 'antes', 'data': HOJE})
    receita_id = consultar('SELECT id FROM receitas')[0][0]

    client.post('/receita/update', data={
        'id': receita_id, 'valor': '250.50', 'categoria': 'Salário', 'descricao': 'depois', 'data': HOJE,
    })

    assert consultar('SELECT descricao, valor FROM receitas')[0] == ('depois', 250.5)


def test_editar_despesa(client):
    entrar(client)
    client.post('/despesas', data={'valor': '80', 'categoria': 'Mercado', 'descricao': 'antes', 'data': HOJE})
    despesa_id = consultar('SELECT id FROM despesas')[0][0]

    client.post('/despesa/update', data={
        'id': despesa_id, 'valor': '95', 'categoria': 'Mercado', 'descricao': 'depois', 'data': HOJE,
    })

    assert consultar('SELECT descricao, valor FROM despesas')[0] == ('depois', 95.0)


def test_valor_invalido_nao_e_salvo(client):
    entrar(client)
    client.post('/receitas', data={'valor': 'abc', 'categoria': 'x', 'descricao': 'x', 'data': HOJE})
    client.post('/despesas', data={'valor': '-10', 'categoria': 'x', 'descricao': 'x', 'data': HOJE})

    assert consultar('SELECT COUNT(*) FROM receitas')[0][0] == 0
    assert consultar('SELECT COUNT(*) FROM despesas')[0][0] == 0


def test_filtro_por_periodo(client):
    entrar(client)
    client.post('/receitas', data={'valor': '100', 'categoria': 'Salário', 'descricao': 'deste mês', 'data': HOJE})
    client.post('/receitas', data={'valor': '100', 'categoria': 'Salário', 'descricao': 'de 2020', 'data': '2020-01-15'})

    padrao = client.get('/receitas').get_data(as_text=True)
    assert 'deste mês' in padrao and 'de 2020' not in padrao

    tudo = client.get('/receitas?mes=&ano=').get_data(as_text=True)
    assert 'deste mês' in tudo and 'de 2020' in tudo

    janeiro_2020 = client.get('/receitas?mes=01&ano=2020').get_data(as_text=True)
    assert 'deste mês' not in janeiro_2020 and 'de 2020' in janeiro_2020


# ---------------------------
# Relatório e formatação
# ---------------------------
def test_relatorio_excel(client):
    entrar(client)
    client.post('/receitas', data={'valor': '100', 'categoria': 'Salário', 'descricao': 'teste', 'data': HOJE})

    resposta = client.get('/gerar_relatorio')

    assert resposta.status_code == 200
    planilha = openpyxl.load_workbook(BytesIO(resposta.data))
    assert planilha.sheetnames == ['Receitas', 'Despesas', 'Investimentos']
    assert planilha['Receitas'].max_row == 2  # cabeçalho + 1 receita


def test_formatar_brl():
    assert finpro.formatar_brl(1234.5) == '1.234,50'
    assert finpro.formatar_brl(0) == '0,00'
    assert finpro.formatar_brl(None) == '0,00'
    assert finpro.formatar_brl(-3.456) == '-3,46'


def test_formatar_data():
    assert finpro.formatar_data('2025-10-14') == '14/10/2025'
    assert finpro.formatar_data('2025-10-14 08:30:00') == '14/10/2025'


# ---------------------------
# Conta demo
# ---------------------------
def test_conta_demo_e_criada_uma_vez_so(app):
    finpro.criar_conta_demo()
    finpro.criar_conta_demo()

    assert consultar('SELECT COUNT(*) FROM usuarios WHERE username = ?', (finpro.DEMO_USUARIO,))[0][0] == 1
    assert consultar('SELECT COUNT(*) FROM receitas')[0][0] > 0
    assert consultar('SELECT COUNT(*) FROM investimentos')[0][0] > 0
    assert consultar('SELECT COUNT(*) FROM despesas WHERE data > ?', (HOJE,))[0][0] == 0


def test_entrar_com_a_conta_demo(app):
    finpro.criar_conta_demo()
    client = app.test_client()

    resposta = client.post('/login', data={'username': finpro.DEMO_USUARIO, 'senha': finpro.DEMO_SENHA})

    assert resposta.status_code == 302
    assert client.get('/dashboard').status_code == 200
